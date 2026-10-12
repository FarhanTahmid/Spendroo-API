from datetime import date
from importlib import import_module
from decimal import Decimal
from io import BytesIO
from unittest import mock

from django.apps import apps
from openpyxl import load_workbook

from billing.catalog.keys import F
from billing.models import Feature, FunnelEvent, Plan, PlanFeature
from billing.testing import grant_plan, set_mode, set_plan_value
from cashbook.models import CashBookAdditionalMember
from reports import views
from reports.models import ReportDownload
from workspaces.models import Membership

from .base import REPORTS, ReportTestCase

EXCEL = {'file_type': 'xlsx'}


class DownloadTests(ReportTestCase):
    def workbook(self, url, **params):
        response = self.alice_client.get(url, {**EXCEL, **params})
        self.assertEqual(response.status_code, 200, getattr(response, 'data', None))
        return load_workbook(BytesIO(response.content))

    def test_every_report_is_a_pdf_on_the_free_plan(self):
        for scope, url in self.urls().items():
            response = self.alice_client.get(url)
            self.assertEqual(response.status_code, 200, scope)
            self.assertEqual(response['Content-Type'], 'application/pdf')
            self.assertTrue(response.content.startswith(b'%PDF'), scope)
        self.assertEqual(
            sorted(ReportDownload.objects.values_list('scope', flat=True)),
            ['account', 'business', 'cashbook', 'personal'],
        )

    def test_a_download_is_noted_and_carries_its_reference(self):
        response = self.alice_client.get(self.urls()['account'], {'period': 'last_month', 'lang': 'bn'})
        self.assertEqual(response.status_code, 200)
        noted = ReportDownload.objects.get()
        self.assertEqual(response['X-Report-Reference'], noted.reference)
        self.assertRegex(noted.reference, r'^SR-20261012-[2-9A-Z]{6}$')
        self.assertEqual(
            (noted.user, noted.workspace, noted.scope, noted.subject_id, noted.subject_name, noted.file_type),
            (self.alice, self.personal, 'account', self.cash.id, 'Cash', 'pdf'),
        )
        self.assertEqual(
            (noted.period, noted.date_from, noted.date_to, noted.language, noted.row_count),
            ('last_month', date(2026, 9, 1), date(2026, 9, 30), 'bn', 3),
        )
        self.assertEqual(
            response['Content-Disposition'], 'attachment; filename="spendroo-account-cash-20260901-20260930.pdf"',
        )

    def test_excel_is_a_plan_feature(self):
        for scope, url in self.urls().items():
            response = self.alice_client.get(url, EXCEL)
            self.assertEqual(response.status_code, 402, scope)
            self.assertEqual((response.data['code'], response.data['feature']), ('plan_limit', 'reports.excel'))
            self.assertEqual(response.data['upgrade_to'], 'plus')
        self.assertFalse(ReportDownload.objects.exists())
        self.assertEqual(FunnelEvent.objects.filter(feature_key='reports.excel').count(), 4)

        grant_plan(self.alice, 'plus')
        for scope, url in self.urls().items():
            self.assertEqual(self.alice_client.get(url, EXCEL).status_code, 200, scope)

    def test_which_plans_get_excel_is_decided_in_the_database(self):
        set_plan_value('free', F.REPORTS_EXCEL, enabled=True)
        self.assertEqual(self.alice_client.get(self.urls()['cashbook'], EXCEL).status_code, 200)

    def test_nothing_is_refused_while_enforcement_is_off(self):
        set_mode('log_only')
        self.assertEqual(self.alice_client.get(self.urls()['cashbook'], EXCEL).status_code, 200)

    def test_chosen_dates_are_a_plan_feature_and_presets_are_not(self):
        custom = {'period': 'custom', 'from': '2026-09-10', 'to': '2026-10-02'}
        url = self.urls()['account']
        for preset in ('this_month', 'last_month', 'last_3_months', 'this_year'):
            self.assertEqual(self.alice_client.get(url, {'period': preset}).status_code, 200, preset)
        response = self.alice_client.get(url, custom)
        self.assertEqual((response.status_code, response.data['feature']), (402, 'reports.custom_range'))
        grant_plan(self.alice, 'plus')
        self.assertEqual(self.alice_client.get(url, custom).status_code, 200)
        noted = ReportDownload.objects.latest('created_at')
        self.assertEqual((noted.period, noted.date_from, noted.date_to), ('custom', date(2026, 9, 10), date(2026, 10, 2)))

    def test_the_workbook_holds_the_same_figures_as_real_numbers(self):
        grant_plan(self.alice, 'plus')
        book = self.workbook(self.urls()['account'])
        self.assertEqual(book.sheetnames, ['Summary', 'Breakdown', 'Details'])
        summary = {row[0]: row[1] for row in book['Summary'].iter_rows(values_only=True) if row[0]}
        self.assertEqual(summary['Opening balance'], Decimal('29800'))
        self.assertEqual(summary['Money in'], 5000)
        self.assertEqual(summary['Money out'], 750)
        self.assertEqual(summary['Closing balance'], 34050)
        self.assertEqual(summary['Period'], '1 Oct 2026 to 31 Oct 2026')

        details = list(book['Details'].iter_rows(values_only=True))
        self.assertEqual(details[0][0], 'Statement')
        self.assertEqual(details[2], ('Date', 'Note', 'Category', 'Money in', 'Money out', 'Balance'))
        self.assertEqual(details[4][1:], ('Lunch 50% off', 'Food', None, 450, 29350))
        self.assertEqual(details[4][0].date(), date(2026, 10, 2))
        self.assertEqual(details[-1][1:3] + details[-1][5:], ('Closing balance', None, 34050))
        self.assertIn('SR-20261012-', book['Details'].oddFooter.left.text)

    def test_what_a_user_typed_is_never_a_formula(self):
        grant_plan(self.alice, 'plus')
        self.add('expense', 100, '2026-10-07', note='=HYPERLINK("http://example.com")')
        cells = [cell for row in self.workbook(self.urls()['account'])['Details'].iter_rows() for cell in row]
        typed = next(cell for cell in cells if str(cell.value).startswith('=HYPERLINK'))
        self.assertEqual(typed.data_type, 's')

    def test_in_bangla(self):
        grant_plan(self.alice, 'plus')
        book = self.workbook(self.urls()['business'], lang='bn')
        self.assertEqual(book.sheetnames, ['সারসংক্ষেপ', 'বিশ্লেষণ', 'বিস্তারিত'])
        self.assertEqual(self.alice_client.get(self.urls()['business'], {'lang': 'bn'}).status_code, 200)

    def test_bad_requests_say_what_is_wrong(self):
        url = self.urls()['personal']
        for params in (
            {'file_type': 'csv'}, {'lang': 'fr'}, {'period': 'yesterday'}, {'period': 'custom'},
            {'period': 'custom', 'from': '2026-10-05', 'to': '2026-10-01'},
            {'period': 'custom', 'from': 'soon', 'to': '2026-10-01'},
        ):
            self.assertEqual(self.alice_client.get(url, params).status_code, 400, params)
        self.assertFalse(ReportDownload.objects.exists())

    def test_a_period_with_too_many_lines_is_refused_with_the_numbers(self):
        with mock.patch.dict(views.MAX_ROWS, {'pdf': 2}):
            response = self.alice_client.get(self.urls()['account'])
        self.assertEqual(response.status_code, 400)
        self.assertEqual(
            (response.data['code'], response.data['rows'], response.data['limit']), ('report_too_large', 3, 2),
        )


class AccessTests(ReportTestCase):
    def test_a_stranger_finds_nothing_and_learns_nothing_about_the_plan(self):
        urls = self.urls()
        for scope in ('account', 'cashbook', 'business'):
            for params in ({}, EXCEL):
                self.assertEqual(self.bob_client.get(urls[scope], params).status_code, 404, scope)
        self.assertEqual(self.bob_client.get(f'{REPORTS}businesses/{self.personal.id}/').status_code, 404)

    def test_the_personal_report_is_always_ones_own(self):
        self.assertEqual(self.bob_client.get(self.urls()['personal']).status_code, 200)
        noted = ReportDownload.objects.get()
        self.assertEqual((noted.user, noted.workspace.owner), (self.bob, self.bob))

    def test_staff_get_the_books_they_were_given_and_not_the_whole_business(self):
        urls = self.urls()
        Membership.objects.create(workspace=self.shop, user=self.bob, role=Membership.ROLE_STAFF)
        self.assertEqual(self.bob_client.get(urls['business']).status_code, 404)
        self.assertEqual(self.bob_client.get(urls['cashbook']).status_code, 404)
        CashBookAdditionalMember.objects.create(cashbook=self.till, member=self.bob, role='viewer')
        self.assertEqual(self.bob_client.get(urls['cashbook']).status_code, 200)
        self.assertEqual(self.bob_client.get(urls['business']).status_code, 404)

    def test_a_viewer_gets_the_whole_business_on_the_owners_plan(self):
        Membership.objects.create(workspace=self.shop, user=self.bob, role=Membership.ROLE_VIEWER)
        url = self.urls()['business']
        self.assertEqual(self.bob_client.get(url).status_code, 200)
        self.assertEqual(self.bob_client.get(url, EXCEL).status_code, 402)
        grant_plan(self.alice, 'business')
        self.assertEqual(self.bob_client.get(url, EXCEL).status_code, 200)

    def test_signed_out_gets_nothing(self):
        from rest_framework.test import APIClient
        self.assertIn(APIClient().get(self.urls()['personal']).status_code, (401, 403))


class FlagRenameTests(ReportTestCase):
    """billing/migrations/0003: the two cashbook flags became the two report flags."""

    def test_a_plan_keeps_what_the_admin_gave_it(self):
        Feature.objects.filter(key__startswith='reports.').delete()
        old = Feature.objects.create(key='cashbook.report_export', kind='flag', name='Old', is_system=True)
        PlanFeature.objects.create(plan=Plan.objects.get(code='free'), feature=old, enabled=True)

        import_module('billing.migrations.0003_report_flags').rename_report_flags(apps, None)

        self.assertFalse(Feature.objects.filter(key__startswith='cashbook.report').exists())
        renamed = Feature.objects.get(key='reports.excel')
        self.assertEqual((renamed.pk, renamed.name), (old.pk, 'Reports as Excel'))
        # The admin had switched it on for Free; the default would have been off
        self.assertTrue(PlanFeature.objects.get(plan__code='free', feature=renamed).enabled)
        self.assertTrue(PlanFeature.objects.get(plan__code='plus', feature__key='reports.custom_range').enabled)
