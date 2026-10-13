"""A custom report narrowed to kinds of line and categories."""
from io import BytesIO

from openpyxl import load_workbook

from billing.testing import grant_plan
from cashbook.models import EntryCategory
from ledger_personal.models import Category
from reports.builders import account, business, cashbook, personal
from reports.document import Money, Totals, net
from reports.models import ReportDownload
from reports.narrow import Narrow

from .base import ReportTestCase


def taka(amount):
    return Money(amount, 'BDT')


class NarrowedBuilderTests(ReportTestCase):
    def test_an_account_narrowed_to_spending_shows_no_balances(self):
        report = account.build(self.cash, self.context(narrow=Narrow(types=('expense',))))
        self.assertEqual(report.totals, [net('BDT', 0, 75000)])
        self.assertEqual(report.totals[0].closing_minor, -75000)
        self.assertEqual(report.filters, [('Type', 'Spending')])
        self.assertIn('Balances are left out, because the lines not shown would be missing from them.', report.notes)

        statement = self.table(report, 'Statement')
        self.assertEqual([column.label for column in statement.columns], ['Date', 'Note', 'Category', 'Money in', 'Money out'])
        self.assertIsNone(statement.first)
        self.assertEqual([row[1] for row in statement.rows], ['Lunch 50% off', 'Rickshaw'])
        self.assertEqual(statement.total, [None, 'Total', '', taka(0), taka(75000)])
        self.assertEqual([table.title for table in report.tables], ['What moved', 'Spending by category', 'Statement'])
        self.assertEqual([row[0] for row in self.table(report, 'What moved').rows], ['Spending'])

    def test_an_account_narrowed_to_a_category(self):
        narrow = Narrow(category_ids=(self.food.id,), category_names=('Food',))
        report = account.build(self.cash, self.context(narrow=narrow))
        self.assertEqual(report.filters, [('Categories', 'Food')])
        self.assertEqual([row[1] for row in self.table(report, 'Statement').rows], ['Lunch 50% off'])
        self.assertEqual(report.totals, [net('BDT', 0, 45000)])

    def test_a_split_narrowed_to_a_category_counts_for_that_part(self):
        transport = Category.objects.get(workspace=self.personal, name='Transport')
        self.split('expense', '2026-10-06', [(self.food, 20000), (transport, 10000), (None, 20000)], note='ATM cash')
        narrow = Narrow(category_ids=(self.food.id,), uncategorised=True, category_names=('Food',))
        report = account.build(self.cash, self.context(narrow=narrow))
        self.assertEqual(report.totals, [net('BDT', 0, 115000)])
        statement = self.table(report, 'Statement')
        self.assertEqual(statement.rows[-1][1:], ['ATM cash', 'Split: Food, Uncategorised', None, taka(40000)])
        self.assertEqual(statement.total, [None, 'Total', '', taka(0), taka(115000)])
        self.assertEqual(self.table(report, 'Spending by category').total[2], taka(115000))

        whole = personal.build(self.personal, self.context(narrow=Narrow(category_ids=(transport.id,))))
        self.assertEqual(whole.totals, [net('BDT', 0, 10000)])
        self.assertEqual(self.table(whole, 'Income and spending').rows, [
            [statement.rows[-1][0], 'Cash', 'Split: Transport', 'ATM cash', None, taka(10000)],
        ])

    def test_uncategorised_means_income_and_spending_with_no_category(self):
        report = account.build(self.cash, self.context(narrow=Narrow(uncategorised=True)))
        self.assertEqual(report.filters, [('Categories', 'Uncategorised')])
        # The transfer has no category either, and is not what was asked for
        self.assertEqual([row[1] for row in self.table(report, 'Statement').rows], ['Rickshaw'])

    def test_everything_still_has_its_balances(self):
        report = account.build(self.cash, self.context())
        self.assertEqual(report.filters, [])
        self.assertEqual(report.totals, [Totals('BDT', 2980000, 500000, 75000, 3405000)])

    def test_the_personal_report_narrowed(self):
        narrow = Narrow(types=('expense',), category_ids=(self.food.id,), uncategorised=True, category_names=('Food',))
        report = personal.build(self.personal, self.context(narrow=narrow))
        self.assertEqual(report.totals, [net('BDT', 0, 75000)])
        self.assertEqual(report.filters, [('Type', 'Spending'), ('Categories', 'Food, Uncategorised')])
        self.assertEqual(
            [table.title for table in report.tables], ['Spending by category', 'Income and spending'],
        )
        self.assertIn('Net = income − spending, for the lines shown.', report.notes)

    def test_the_personal_report_narrowed_to_transfers_lists_only_them(self):
        report = personal.build(self.personal, self.context(narrow=Narrow(types=('transfer',))))
        self.assertEqual([table.title for table in report.tables], ['Transfers between your accounts'])
        self.assertEqual(report.row_count, 1)

    def test_a_cashbook_narrowed_to_cash_out(self):
        report = cashbook.build(self.till, self.context(narrow=Narrow(types=('cash_out',))))
        self.assertEqual(report.totals, [net('BDT', 0, 20000)])
        self.assertEqual(report.filters, [('Type', 'Cash out')])
        entries = self.table(report, 'Entries')
        self.assertNotIn('Balance', [column.label for column in entries.columns])
        self.assertEqual([row[1] for row in entries.rows], ['Shop rent'])
        self.assertEqual(entries.total[-2:], [None, taka(20000)])
        self.assertEqual(self.table(report, 'By category').rows, [['Rent', 1, None, taka(20000)]])

    def test_a_business_narrowed_to_a_category_found_in_its_cashbooks(self):
        other = EntryCategory.objects.create(cashbook=self.safe, category_name='Sales')
        self.entry(self.safe, 'cash_in', 3000, '2026-10-06', title='Safe sale', category=other)
        narrow = Narrow(category_ids=(self.sales.id, other.id), category_names=('Sales',))
        report = business.build(self.shop, self.context(narrow=narrow))
        self.assertEqual(report.totals, [net('BDT', 53000, 0)])
        books = self.table(report, 'Cashbooks')
        self.assertEqual([column.label for column in books.columns], ['Cashbook', 'Money in', 'Money out', 'Net'])
        self.assertEqual(books.rows, [['Till', taka(50000), None, taka(50000)], ['Safe', taka(3000), None, taka(3000)]])
        self.assertEqual(books.total, ['Total', taka(53000), taka(0), taka(53000)])

    def test_in_bangla(self):
        report = cashbook.build(self.till, self.context(language='bn', narrow=Narrow(types=('cash_in',))))
        self.assertEqual(report.filters, [('ধরন', 'ক্যাশ ইন')])


class NarrowedDownloadTests(ReportTestCase):
    CUSTOM = {'period': 'custom', 'from': '2026-09-01', 'to': '2026-10-31'}

    def test_a_filter_is_part_of_the_custom_plan_feature_whatever_the_period(self):
        for params in ({'type': 'expense'}, {'category': 'none'}, {'category': str(self.food.id), 'period': 'last_month'}):
            response = self.alice_client.get(self.urls()['account'], params)
            self.assertEqual((response.status_code, response.data['feature']), (402, 'reports.custom_range'), params)
        self.assertFalse(ReportDownload.objects.exists())

    def test_a_custom_report_carries_its_dates_types_and_categories(self):
        grant_plan(self.alice, 'plus')
        params = {**self.CUSTOM, 'type': 'expense,income', 'category': f'{self.food.id},none', 'file_type': 'xlsx'}
        response = self.alice_client.get(self.urls()['account'], params)
        self.assertEqual(response.status_code, 200, getattr(response, 'data', None))

        noted = ReportDownload.objects.get()
        self.assertEqual(noted.filters, {'Type': 'Income, Spending', 'Categories': 'Food, Uncategorised'})
        self.assertEqual(noted.row_count, 3)

        summary = {row[0]: row[1] for row in load_workbook(BytesIO(response.content))['Summary'].iter_rows(values_only=True) if row[0]}
        self.assertEqual(summary['Type'], 'Income, Spending')
        self.assertEqual(summary['Categories'], 'Food, Uncategorised')
        self.assertEqual(summary['Money out'], 1950)
        self.assertEqual(summary['Net (in − out)'], -1950)
        self.assertNotIn('Opening balance', summary)

    def test_every_report_takes_a_filter_and_renders_as_a_pdf(self):
        grant_plan(self.alice, 'plus')
        urls = self.urls()
        for scope, params in (
            ('personal', {'type': 'expense'}), ('account', {'category': 'none'}),
            ('cashbook', {'type': 'cash_in', 'category': str(self.sales.id)}), ('business', {'type': 'cash_out'}),
        ):
            response = self.alice_client.get(urls[scope], {**params, 'lang': 'bn'})
            self.assertEqual(response.status_code, 200, scope)
            self.assertTrue(response.content.startswith(b'%PDF'), scope)

    def test_every_kind_chosen_is_no_filter(self):
        response = self.alice_client.get(self.urls()['cashbook'], {'type': 'cash_in,cash_out'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(ReportDownload.objects.get().filters, {})

    def test_a_report_only_takes_its_own_kinds_and_categories(self):
        grant_plan(self.alice, 'plus')
        urls = self.urls()
        for scope, params in (
            ('account', {'type': 'cash_in'}), ('cashbook', {'type': 'expense'}),
            ('account', {'category': 'food'}),
            # A cashbook's category on a personal report, and the other way round
            ('account', {'category': str(self.sales.id)}), ('cashbook', {'category': str(self.food.id)}),
            # A category of the business, but of another cashbook
            (None, {'category': str(self.sales.id)}),
        ):
            url = urls[scope] if scope else f'/api/v1/reports/cashbooks/{self.safe.id}/'
            self.assertEqual(self.alice_client.get(url, params).status_code, 400, (scope, params))

    def test_a_filter_cannot_reach_into_someone_elses_categories(self):
        grant_plan(self.bob, 'plus')
        response = self.bob_client.get(self.urls()['personal'], {'category': str(self.food.id)})
        self.assertEqual(response.status_code, 400)
