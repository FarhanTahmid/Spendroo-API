"""
A Report as an Excel workbook: Summary, Breakdown, Details.

Amounts are real numbers and dates are real dates, so the sheets can be
sorted, filtered and added up. The words are the same ones the PDF prints.
"""
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from spendroo.money import currency_exponent, to_major

from decimal import Decimal

from ..document import COUNT, DATE, MONEY, NUMBER, SHARE, TONE_BALANCE, TONE_IN, TONE_OUT
from ..labels import Words
from . import cells

ASSETS = Path(__file__).resolve().parent.parent / 'assets'

NAVY, GOLD, GREEN, RED, MUTED = '0B2D5E', 'F5B530', '17804A', 'D93F4C', '5B6B84'
HEAD_FILL = PatternFill('solid', start_color=NAVY, end_color=NAVY)
EDGE_FILL = PatternFill('solid', start_color='EEF2F8', end_color='EEF2F8')
NOTE_FILL = PatternFill('solid', start_color='FFF8E6', end_color='FFF8E6')
HAIRLINE = Border(bottom=Side(style='thin', color='E2E8F1'))
TOTAL_LINE = Border(top=Side(style='medium', color=NAVY))
TONE_COLOR = {TONE_IN: GREEN, TONE_OUT: RED}
# Characters of width one unit of Column.width is worth
UNIT_WIDTH = 9
LOGO_HEIGHT = 46


def money_format(currency, named):
    """'#,##0.00', with the code in front when the table does not name its currency."""
    digits = currency_exponent(currency)
    number = '#,##0' + ('.' + '0' * digits if digits else '')
    # The sign goes next to the number: "BDT -120.00", not "-BDT 120.00"
    return number if named else f'"{currency}" {number};"{currency}" -{number}'


def _write(sheet, row, column, value, *, font=None, fill=None, border=None, align=None, number_format=None):
    cell = sheet.cell(row=row, column=column, value=value)
    if isinstance(value, str):
        # What a user typed is text, never a formula, even when it starts with "="
        cell.data_type = 's'
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if border:
        cell.border = border
    if align:
        cell.alignment = align
    if number_format:
        cell.number_format = number_format
    return cell


def _cell(sheet, row, index, column, value, table, *, bold=False, fill=None, border=HAIRLINE):
    """One table cell, as a value Excel can work with."""
    font = Font(bold=bold or column.tone == TONE_BALANCE, color=TONE_COLOR.get(column.tone))
    align = Alignment(vertical='top', wrap_text=column.kind not in (MONEY, COUNT, SHARE, DATE, NUMBER))
    number_format = None
    if value is None or value == '':
        value = None
    elif column.kind == MONEY:
        amount_minor, currency = value
        value, number_format = to_major(amount_minor, currency), money_format(currency, bool(table.currency))
    elif column.kind == DATE:
        number_format = 'DD MMM YYYY'
    elif column.kind == COUNT:
        number_format = '#,##0'
    elif column.kind == SHARE:
        number_format = '0.0%'
    elif column.kind == NUMBER and isinstance(value, Decimal):
        # As many decimals as were typed: 12.5 stays 12.5, 1200 stays 1,200
        places = max(0, -value.as_tuple().exponent)
        number_format = '#,##0' + ('.' + '0' * places if places else '')
    _write(sheet, row, index, value, font=font, fill=fill, border=border, align=align, number_format=number_format)


def _table(sheet, row, table, words):
    """Writes the table from `row` down. Returns (the row of its headings, the next free row)."""
    _write(sheet, row, 1, table.title, font=Font(bold=True, size=12, color=NAVY))
    row += 1
    for line in (words.t('table.amounts_in', currency=table.currency) if table.currency else '', table.note):
        if line:
            _write(sheet, row, 1, line, font=Font(color=MUTED, italic=True))
            row += 1
    header = row
    for index, column in enumerate(table.columns, 1):
        numeric = column.kind in (MONEY, COUNT, SHARE, NUMBER)
        _write(
            sheet, row, index, column.label, font=Font(bold=True, color='FFFFFF'), fill=HEAD_FILL,
            align=Alignment(horizontal='right' if numeric else 'left', vertical='center'),
        )
    row += 1
    if table.first:
        for index, (column, value) in enumerate(zip(table.columns, table.first), 1):
            _cell(sheet, row, index, column, value, table, bold=True, fill=EDGE_FILL)
        row += 1
    for values in table.rows:
        for index, (column, value) in enumerate(zip(table.columns, values), 1):
            _cell(sheet, row, index, column, value, table)
        row += 1
    if not table.rows and table.empty:
        _write(sheet, row, 1, table.empty, font=Font(color=MUTED, italic=True))
        row += 1
    if table.total:
        for index, (column, value) in enumerate(zip(table.columns, table.total), 1):
            _cell(sheet, row, index, column, value, table, bold=True, fill=EDGE_FILL, border=TOTAL_LINE)
        row += 1
    return header, row + 1


def _widen(sheet, tables):
    """Each column as wide as the widest table using it needs."""
    widths = {}
    for table in tables:
        for index, column in enumerate(table.columns, 1):
            widths[index] = max(widths.get(index, 0), column.width * UNIT_WIDTH)
    for index, width in widths.items():
        sheet.column_dimensions[get_column_letter(index)].width = width


def _furnish(sheet, report, words):
    """What every printed page of every sheet carries."""
    sheet.sheet_view.showGridLines = False
    # "&" starts a code in a print header
    sheet.oddHeader.left.text = f'Spendroo  ·  {report.title}: {report.subject}'.replace('&', '&&')
    sheet.oddFooter.left.text = f"{words.t('footer.by')}  ·  {report.reference}"
    sheet.oddFooter.right.text = f"{words.t('footer.page')} &P {words.t('footer.of')} &N"
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _stack(sheet, tables, report, words):
    _furnish(sheet, report, words)
    row = 1
    for table in tables:
        header, row = _table(sheet, row, table, words)
        if len(tables) == 1:
            sheet.freeze_panes = sheet.cell(row=header + 1, column=1)
    _widen(sheet, tables)


def _summary(sheet, report, words):
    _furnish(sheet, report, words)
    sheet.column_dimensions['A'].width = 30
    sheet.column_dimensions['B'].width = 46

    logo = Image(ASSETS / 'spendroo-logo-horizontal-800w.png')
    logo.width, logo.height = round(LOGO_HEIGHT * logo.width / logo.height), LOGO_HEIGHT
    sheet.add_image(logo, 'A1')
    for row in (1, 2, 3):
        sheet.row_dimensions[row].height = 18

    _write(sheet, 5, 1, report.title.upper(), font=Font(bold=True, size=10, color=GOLD))
    _write(sheet, 6, 1, report.subject, font=Font(bold=True, size=18, color=NAVY))
    _write(sheet, 7, 1, report.covers, font=Font(color=MUTED))

    row = 9
    for label, value in (
        ('head.period', words.period(report.date_from, report.date_to)),
        ('head.prepared_for', report.prepared_for),
        ('head.generated', words.moment(report.generated_at)),
        ('head.reference', report.reference),
    ):
        _write(sheet, row, 1, words.t(label), font=Font(color=MUTED))
        _write(sheet, row, 2, value, font=Font(bold=True))
        row += 1

    if report.filters:
        row += 1
        _write(sheet, row, 1, words.t('filter.title'), font=Font(bold=True, size=12, color=NAVY))
        _write(sheet, row, 2, words.t('filter.about'), font=Font(color=MUTED, italic=True))
        row += 1
        for label, chosen in report.filters:
            _write(sheet, row, 1, label, font=Font(color=MUTED), fill=NOTE_FILL)
            _write(sheet, row, 2, chosen, font=Font(bold=True), fill=NOTE_FILL, align=Alignment(wrap_text=True))
            row += 1

    row += 1
    _write(sheet, row, 1, words.t('section.glance'), font=Font(bold=True, size=12, color=NAVY))
    row += 1
    for totals in report.totals:
        for kind, label, amount in cells.tiles(words, report, totals):
            closing = kind == 'closing'
            _write(
                sheet, row, 1, label, font=Font(bold=closing, color='FFFFFF' if closing else None),
                fill=HEAD_FILL if closing else None, border=HAIRLINE,
            )
            _write(
                sheet, row, 2, to_major(amount, totals.currency), border=HAIRLINE,
                font=Font(bold=True, color='FFFFFF' if closing else {'in': GREEN, 'out': RED}.get(kind)),
                fill=HEAD_FILL if closing else None, number_format=money_format(totals.currency, False),
                align=Alignment(horizontal='right'),
            )
            row += 1
        row += 1

    _write(sheet, row, 1, words.t('section.how'), font=Font(bold=True, size=12, color=NAVY))
    row += 1
    for note in report.notes:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        _write(sheet, row, 1, note, fill=NOTE_FILL, align=Alignment(wrap_text=True, vertical='top'))
        # A merged cell does not grow with its text
        sheet.row_dimensions[row].height = 15 * (1 + len(note) // 70)
        row += 1


def render(report):
    """The report as .xlsx bytes."""
    words = Words(report.language)
    book = Workbook()
    summary = book.active
    summary.title = words.t('sheet.summary')
    _summary(summary, report, words)
    _stack(book.create_sheet(words.t('sheet.breakdown')), report.breakdowns, report, words)
    _stack(book.create_sheet(words.t('sheet.details')), report.details, report, words)
    book.properties.creator = 'Spendroo'
    book.properties.title = f'{report.title}: {report.subject}'
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
